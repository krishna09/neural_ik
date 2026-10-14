#!/usr/bin/env python
"""Write presentation/speaker_notes.md from the speaker notes stored in the deck itself.

    python deliverables/scripts/deck/export_notes.py deliverables/presentation/Neural_IK_STP_Team06_Level1.pptx
"""
import sys
from pathlib import Path

from pptx import Presentation


def main(path):
    prs = Presentation(path)
    out = ["# Speaker notes — Neural Kinematics for Intelligent Space Manipulators (STP Team 06)", "",
           "Generated from the notes inside `Neural_IK_STP_Team06_Level1.pptx` "
           "(scripts/deck/build_deck.js). About 18–20 minutes in total; timings per slide are suggestions.", ""]
    for i, slide in enumerate(prs.slides, 1):
        title = next((sh.text_frame.text for sh in slide.placeholders if sh.placeholder_format.type is not None
                      and "TITLE" in str(sh.placeholder_format.type)), None)
        if not title:                                   # title slide: the title is a named text box
            title = next((sh.text_frame.text for sh in slide.shapes if sh.name == "t_title" and sh.has_text_frame), None)
        if not title:
            texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            title = max(texts[:4], key=len) if texts else f"Slide {i}"
        notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        out += [f"## Slide {i} · {title.strip()}", ""]
        out += [line if line.strip() else "" for line in notes.splitlines()] + [""]
    dest = Path(path).with_name("speaker_notes.md")
    dest.write_text("\n".join(out), encoding="utf-8")
    print("wrote", dest)


if __name__ == "__main__":
    main(sys.argv[1])

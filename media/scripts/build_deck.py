#!/usr/bin/env python
"""Build the 14-slide deck (brief section 8) -> media/out/neural_ik_deck.pptx.

python-pptx, 16:9, the Studio's dark theme; charts are native PowerPoint charts (editable), filled from the run
files through anim_charts.load_facts(); screenshots come from capture_studio.py. Speaker notes on every slide.

    python media/scripts/build_deck.py
    # render to PNG for checking (Windows PowerPoint): media/scripts/render_deck.ps1
"""
import math
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_AXIS_CROSSES, XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import anim_charts as A  # noqa: E402

ROOT = A.ROOT
FRAMES = ROOT / "media" / "assets" / "frames"
IMG = ROOT / "media" / "assets" / "deck"
OUT = ROOT / "media" / "out" / "neural_ik_deck.pptx"

FONT, FONT_SB = "Segoe UI", "Segoe UI Semibold"
FOOTER = "STP 2026 · Team 06 · simulation results"
DESIGN_ITER = [(30, 70), (37, 55), (67, 3.5), (78, 2.5), (84, 2.0)]   # fact sheet section 2 (no data file holds them)
SW, SH = 13.333, 7.5
M = 0.6                                                                # side margin (in)


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ================================================================================================ images
def prep_images(F):
    IMG.mkdir(parents=True, exist_ok=True)
    out = {}
    # logo: conic gradient in a rounded square, as in the Studio
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyBboxPatch
    fig = plt.figure(figsize=(5.12, 5.12), dpi=100)
    fig.patch.set_alpha(0)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 512)
    ax.set_ylim(0, 512)
    ax.axis("off")
    clip = FancyBboxPatch((8, 8), 496, 496, boxstyle="round,pad=0,rounding_size=158", fc="none", ec="none", transform=ax.transData)
    ax.add_patch(clip)
    import numpy as np
    im = ax.imshow(np.dstack([A.LOGO_RGB, np.ones(A.LOGO_ANG.shape)]), extent=(8, 504, 8, 504), origin="upper", interpolation="bilinear")
    im.set_clip_path(clip)
    out["logo"] = IMG / "logo.png"
    fig.savefig(out["logo"], transparent=True)
    plt.close(fig)

    def crop_robot(src, dst, pad=40):
        im = Image.open(src).convert("RGB")
        w, h = im.size
        box = (0, 60, w, h - 120)                                     # skip top bar and bottom tabs
        a = np.asarray(im.crop(box)).astype(int)
        mask = (a.max(-1) > 110) & ((a.max(-1) - a.min(-1) > 40) | (a.min(-1) > 150))
        ys, xs = np.nonzero(mask)
        x0, x1 = max(0, xs.min() - pad), min(a.shape[1], xs.max() + pad)
        y0, y1 = max(0, ys.min() - pad), min(a.shape[0], ys.max() + pad)
        from PIL import ImageFilter
        alpha = Image.fromarray((mask * 255).astype("uint8")).filter(ImageFilter.MaxFilter(5)).filter(ImageFilter.GaussianBlur(1.2))
        rgba = im.crop(box).convert("RGBA")
        rgba.putalpha(alpha)
        rgba.crop((x0, y0, x1, y1)).save(dst)
        return dst

    for name in ("xarm6", "kuka_iiwa", "panda"):
        src = FRAMES / f"robot_{name}" / "f_00000.jpg"
        if src.exists():
            out[f"robot_{name}"] = crop_robot(src, IMG / f"robot_{name}.png")

    def crop(src, dst, box):
        if src.exists():
            Image.open(src).crop(box).save(dst)
            out[dst.stem] = dst

    crop(FRAMES / "browse" / "f_00000.jpg", IMG / "browse.png", (360, 52, 1920, 52 + 878))
    run = sorted((FRAMES / "mission_run").glob("f_*.jpg"))
    if run:
        pick = run[min(len(run) - 1, int(len(run) * 0.42))]           # agent bending round, straight path already collided
        crop(pick, IMG / "steer.png", (0, 52, 1920, 52 + 1012))
    plan = FRAMES / "mission_plan" / "f_00000.jpg"
    crop(plan, IMG / "mission_ui.png", (0, 0, 1920, 1080))
    return out


# ================================================================================================ drawing helpers
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SW), Inches(SH)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def slide(self, title=None, chapter=None, color=None, sub=None, notes=""):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(A.BG)
        if chapter:
            pill(s, M, 0.42, chapter, color)
        if title:
            text(s, M, 0.92, SW - 2 * M, 0.75, title, 32, A.TXT, bold=True, name="title")
        if sub:
            text(s, M, 1.62, SW - 2 * M, 0.45, sub, 16, A.TXT2, name="subtitle")
        text(s, M, SH - 0.5, 7, 0.3, FOOTER, 10, A.MUTED, name="footer")
        slide_number(s, self.n)
        s.notes_slide.notes_text_frame.text = notes
        return s


def _no_shadow(shape):
    shape.shadow.inherit = False


def text(s, x, y, w, h, content, size=16, color=A.TXT, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         name=None, spacing_after=0, font=None):
    """content: str, or list of paragraphs; a paragraph is str or list of (text, {size,color,bold}) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing_after:
            p.space_after = Pt(spacing_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for t, o in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            b = o.get("bold", bold)
            f.name = o.get("font", font or (FONT_SB if b else FONT))
            f.size = Pt(o.get("size", size))
            f.color.rgb = rgb(o.get("color", color))
    return tb


def box(s, x, y, w, h, fill=A.SURF, line=A.BORDER, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    _no_shadow(sh)
    return sh


def dot(s, cx, cy, r, color, label=None, label_color="FFFFFF", size=12):
    c = box(s, cx - r, cy - r, 2 * r, 2 * r, fill=color, line=None, shape=MSO_SHAPE.OVAL)
    if label:
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r_ = p.add_run()
        r_.text = label
        r_.font.size, r_.font.name, r_.font.color.rgb = Pt(size), FONT_SB, rgb(label_color)
    return c


def pill(s, x, y, label, color):
    w = 0.42 + 0.105 * len(label)
    box(s, x, y, w, 0.36, fill=A.SURF2, line=A.BORDER, radius=0.5, name="chapter")
    dot(s, x + 0.2, y + 0.18, 0.065, color)
    text(s, x + 0.36, y + 0.065, w - 0.4, 0.25, label, 11, A.TXT, bold=True)


def slide_number(s, n):
    tb = s.shapes.add_textbox(Inches(SW - M - 1), Inches(SH - 0.5), Inches(1), Inches(0.3))
    tb.name = "slide number"
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    fld = etree.SubElement(p._p, qn("a:fld"))
    fld.set("id", "{5C7A2A8E-3A1B-4C55-9E57-0D9F7C1B2A11}")
    fld.set("type", "slidenum")
    rpr = etree.SubElement(fld, qn("a:rPr"))
    rpr.set("lang", "en-US")
    rpr.set("sz", "1000")
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:srgbClr")).set("val", A.MUTED.lstrip("#"))
    etree.SubElement(rpr, qn("a:latin")).set("typeface", FONT)
    etree.SubElement(fld, qn("a:t")).text = str(n)


def picture(s, path, x, y, w=None, h=None, frame=True):
    if frame:
        pad = 0.06
        box(s, x - pad, y - pad, (w or 0) + 2 * pad, (h or 0) + 2 * pad, fill=A.SURF, line=A.BORDER, radius=0.03)
    return s.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def stat(s, x, y, w, value, label, value_size=40, sub=None, accent=None):
    """Big number with a small label under it, on a card."""
    h = 1.55 if sub else 1.3
    box(s, x, y, w, h, fill=A.SURF)
    if accent:
        dot(s, x + 0.3, y + 0.32, 0.07, accent)
    text(s, x + 0.25, y + 0.18, w - 0.4, 0.75, value, value_size, A.TXT, bold=True)
    text(s, x + 0.27, y + 0.88, w - 0.45, 0.3, label, 13, A.TXT2)
    if sub:
        text(s, x + 0.27, y + 1.15, w - 0.45, 0.3, sub, 11, A.MUTED)


# ---- charts
def _chart_space_nofill(chart):
    cs = chart._chartSpace
    sppr = cs.find(qn("c:spPr"))
    if sppr is None:
        sppr = etree.SubElement(cs, qn("c:spPr"))
        cs.remove(sppr)
        txpr = cs.find(qn("c:txPr"))
        (txpr.addprevious(sppr) if txpr is not None else cs.find(qn("c:chart")).addnext(sppr))
    for c in list(sppr):
        sppr.remove(c)
    etree.SubElement(sppr, qn("a:noFill"))
    ln = etree.SubElement(sppr, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))


def style_chart(chart, legend=True, size=12):
    _chart_space_nofill(chart)
    chart.has_title = False
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(A.TXT2)
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
        chart.legend.font.color.rgb = rgb(A.TXT2)
    va, ca = chart.value_axis, chart.category_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb(A.BORDER)
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.tick_labels.font.color.rgb = rgb(A.MUTED)
    va.tick_labels.font.size = Pt(size - 1)
    ca.format.line.color.rgb = rgb(A.BORDER)
    ca.tick_labels.font.color.rgb = rgb(A.TXT2)
    ca.tick_labels.font.size = Pt(size)
    ca.has_major_gridlines = False


def log_axis(axis, lo, hi):
    scaling = axis._element.find(qn("c:scaling"))
    lb = etree.Element(qn("c:logBase"))
    lb.set("val", "10")
    scaling.insert(0, lb)
    axis.minimum_scale, axis.maximum_scale = lo, hi


# ================================================================================================ slides
def build(F, img):
    d = Deck()
    ev = F["l2_eval"]
    comb = ev["combined"]
    md = F["m_dyn"]
    pol = md["controllers"]["policy"]

    # 1 ---------------------------------------------------------------------------------------- title
    s = d.slide(notes=(
        "Team 06, STP 2026. One line: tell the arm where to go – it finds the joint angles in milliseconds, then steers "
        "itself there safely while obstacles appear and its own joints slip.\n\nThree beats: Solve (Level 1, neural IK + "
        "LM solver), Steer (Level 2, a tiny DRL agent), Together (the mission pipeline). All results are simulation "
        "(PyBullet), xArm6, unseen data."))
    s.shapes.add_picture(str(img["logo"]), Inches(M), Inches(0.75), Inches(0.75), Inches(0.75))
    text(s, M + 0.95, 0.86, 6, 0.5, "Neural IK Studio · Team 06 · STP 2026", 16, A.TXT2, bold=True)
    text(s, M, 1.85, 7.2, 2.2, "Neural Kinematics & RL for Collision-Free Space Manipulators", 40, A.TXT, bold=True,
         name="title")
    text(s, M, 4.2, 6.6, 0.9, "Tell the arm where to go – it finds the joint angles in milliseconds, then steers itself "
         "there safely.", 17, A.TXT2)
    beats = [("1", "Solve", A.C_L1), ("2", "Steer", A.C_L2), ("3", "Together", A.C_MISSION)]
    for i, (n, lab, col) in enumerate(beats):
        x = M + i * 2.05
        box(s, x, 5.35, 1.85, 0.5, fill=A.SURF2, radius=0.5)
        dot(s, x + 0.3, 5.6, 0.14, col, n, size=11)
        text(s, x + 0.55, 5.45, 1.2, 0.35, lab, 15, A.TXT, bold=True)
    robots = [k for k in ("xarm6", "kuka_iiwa", "panda") if f"robot_{k}" in img]
    names = dict(F["robots"]) if F["robots"] else {}
    labels = {"xarm6": "xArm6", "kuka_iiwa": "KUKA LBR iiwa", "panda": "Franka Panda"}
    x0, cw = 8.15, 1.6
    for i, k in enumerate(robots):
        x = x0 + i * (cw + 0.05)
        box(s, x, 1.0, cw, 4.85, fill=A.SURF, radius=0.06)
        im = Image.open(img[f"robot_{k}"])
        ih = 3.7
        iw = min(cw - 0.2, ih * im.width / im.height)
        ih = iw * im.height / im.width
        s.shapes.add_picture(str(img[f"robot_{k}"]), Inches(x + (cw - iw) / 2), Inches(1.15 + (3.75 - ih) / 2), Inches(iw), Inches(ih))
        text(s, x + 0.1, 5.05, cw - 0.2, 0.3, labels[k], 13, A.TXT, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 0.1, 5.38, cw - 0.2, 0.3, f"{names.get(labels[k], '?')}-DOF", 11, A.MUTED, align=PP_ALIGN.CENTER)

    # 2 ---------------------------------------------------------------------------------------- problem
    s = d.slide("The problem: reach a pose without ever colliding", "PROJECT BRIEF", A.MUTED,
                notes=("Problem statement (Project Brief – Team 6), verbatim:\n\nLevel 1: Use simulated multi-axis space "
                       "manipulator joint data to develop a kinematics-aware AI/ML model that predicts collision-free "
                       "trajectories and computes joint angles to prevent self collision with the satellite body.\n\n"
                       "Level 2: Implement a lightweight Deep Reinforcement Learning (DRL) policy agent optimized for low "
                       "compute overhead that dynamically re-routes the manipulator path in real time when unexpected "
                       "geometric obstacles or hardware position errors occur.\n\nDataset: OpenRAVE / PyBullet Manipulator "
                       "Kinematics & Space Robot Simulation Dataset."))
    cards = [("LEVEL 1", A.C_L1, "Joint angles without self-collision",
              "“…a kinematics-aware AI/ML model that predicts collision-free trajectories and computes joint angles…”"),
             ("LEVEL 2", A.C_L2, "Re-route in real time, on little compute",
              "“…a lightweight DRL policy agent… that dynamically re-routes the manipulator path in real time…”")]
    for i, (tag, col, head, quote) in enumerate(cards):
        x = M + i * 6.17
        box(s, x, 2.15, 5.9, 3.9, fill=A.SURF)
        pill(s, x + 0.35, 2.45, tag, col)
        text(s, x + 0.35, 3.1, 5.2, 0.9, head, 24, A.TXT, bold=True)
        text(s, x + 0.35, 4.15, 5.2, 1.6, quote, 16, A.TXT2)
    text(s, M, 6.3, 12, 0.35, "Unexpected obstacles and hardware position errors happen mid-move.", 14, A.MUTED)

    # 3 ---------------------------------------------------------------------------------------- pipeline
    s = d.slide("Our answer: Solve, then Steer", "PIPELINE", A.C_MISSION,
                "Five steps from a target pose to a safe, exact final pose",
                notes=("1 Target pose: where the gripper must go (position + orientation).\n2 Level 1 solves IK: the "
                       "network predicts joint angles from the pose and the current joints; a Levenberg–Marquardt solver "
                       "refines them to < 1 mm / < 1°, keeping joint limits and rejecting self-collisions.\n3 Path check: "
                       "the straight joint-space path is checked against known obstacles.\n4 Level 2 steers: the DRL agent "
                       "moves the arm, reacting to obstacles that appear or drift and to joint slip.\n5 Final check: a short "
                       "collision-checked final approach puts the gripper onto the IK goal."))
    steps = [("1", "Target pose", "where the gripper must go", A.MUTED),
             ("2", "Solve IK", "network + LM solver", A.C_L1),
             ("3", "Path check", "straight path vs. obstacles", A.MUTED),
             ("4", "Steer", "DRL agent re-routes", A.C_L2),
             ("5", "Final check", "gripper on the target", A.C_MISSION)]
    w, gap = 2.12, 0.35
    for i, (n, head, sub, col) in enumerate(steps):
        x = M + i * (w + gap)
        box(s, x, 2.75, w, 2.6, fill=A.SURF)
        dot(s, x + 0.45, 3.25, 0.25, col, n, size=14)
        text(s, x + 0.25, 3.8, w - 0.4, 0.5, head, 20, A.TXT, bold=True)
        text(s, x + 0.25, 4.35, w - 0.4, 0.8, sub, 14, A.TXT2)
        if i < 4:
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + w + 0.05), Inches(3.9), Inches(gap - 0.1), Inches(0.3))
            a.fill.solid()
            a.fill.fore_color.rgb = rgb(A.BORDER)
            a.line.fill.background()
            _no_shadow(a)
    for i, (lab, col, x0_, x1_) in enumerate((("Level 1", A.C_L1, 1, 1), ("Level 2", A.C_L2, 3, 3))):
        x = M + x0_ * (w + gap)
        text(s, x + 0.25, 5.6, w, 0.3, lab, 12, A.MUTED, bold=True)

    # 4 ---------------------------------------------------------------------------------------- robots
    s = d.slide("Three arms, one pipeline", "ROBOTS", A.C_L1,
                "Each robot is a configuration entry – a space-arm URDF drops in the same way",
                notes=("robot_config.py registers each arm by URDF and actuated joint names; everything downstream "
                       "(data generation, training, IK, Level 2) reads the robot from there. Results in this deck are "
                       "for the xArm6; KUKA and Panda are supported (data, training, inference), not benchmarked here. "
                       "A space-manipulator URDF would be added the same way."))
    from robot_config import ROBOTS
    for i, k in enumerate(("xarm6", "kuka_iiwa", "panda")):
        x = M + i * 4.1
        box(s, x, 2.3, 3.85, 4.0, fill=A.SURF)
        if f"robot_{k}" in img:
            im = Image.open(img[f"robot_{k}"])
            ih = 2.45
            iw = min(3.4, ih * im.width / im.height)
            ih = iw * im.height / im.width
            s.shapes.add_picture(str(img[f"robot_{k}"]), Inches(x + (3.85 - iw) / 2), Inches(2.45 + (2.5 - ih) / 2), Inches(iw), Inches(ih))
        cfg = ROBOTS[k]
        text(s, x + 0.3, 5.05, 3.3, 0.4, labels[k], 20, A.TXT, bold=True)
        text(s, x + 0.3, 5.5, 3.3, 0.6, [[(f"{cfg['dof']}-DOF", {"bold": True, "color": A.TXT}),
                                          (f" · {len(cfg['joint_names'])} revolute joints", {})],
                                         "benchmarked" if k == "xarm6" else "supported"], 13, A.TXT2)

    # 5 ---------------------------------------------------------------------------------------- data
    ds_n = 200000
    s = d.slide("Data from simulation, not from a download", "LEVEL 1 · DATA", A.C_L1,
                notes=("Dataset ds_20260928-163059_xarm6_n200000_s0: 200,000 samples generated in PyBullet. Joint "
                       "configurations are sampled over the joint limits, with extra samples near the limits (where "
                       "networks usually fail); forward kinematics gives the gripper pose, and PyBullet's contact check "
                       "labels self-collisions. In the Studio's Browse view blue points are collision-free samples, red "
                       "points colliding ones."))
    if "browse" in img:
        picture(s, img["browse"], M, 1.75, w=7.9, h=7.9 * 878 / 1560)
    x = M + 8.35
    text(s, x, 1.85, 3.8, 0.8, f"{ds_n:,}", 44, A.TXT, bold=True)
    text(s, x, 2.65, 3.8, 0.4, "simulated samples (xArm6)", 14, A.TXT2)
    box(s, x, 3.35, 3.8, 1.65, fill=A.SURF)
    text(s, x + 0.25, 3.5, 3.4, 1.4, [[("q", {"bold": True}), (" ~ joint limits, extra near the edges", {})],
                                       [("pose", {"bold": True}), (" = FK(q)", {})],
                                       [("label", {"bold": True}), (" = PyBullet self-collision", {})]], 15, A.TXT2,
         spacing_after=6, font="Segoe UI")
    for j, (lab, col) in enumerate((("collision-free", "#3d8bfd"), ("colliding", "#ff4d4d"))):
        dot(s, x + 0.12, 5.45 + 0.4 * j, 0.08, col)
        text(s, x + 0.35, 5.31 + 0.4 * j, 3, 0.3, lab, 13, A.TXT2)

    # 6 ---------------------------------------------------------------------------------------- level 1 chart
    vals = [F["plan_mm"], F["l1_net_mm"], F["l1_lm_mm"]]
    s = d.slide("Level 1: a network guesses, a solver makes it exact", "LEVEL 1 · SOLVE", A.C_L1,
                "Mean position error on unseen poses (xArm6, log scale)",
                notes=(f"Network alone: {vals[1]:.1f} mm mean (25,090 test poses), 5.5× better than the original plan's "
                       f"{vals[0]:.0f} mm. Inputs use a sin/cos encoding of the angles plus a seed configuration (the "
                       "current joints). The LM step: Δθ = (JᵀJ + λI)⁻¹ Jᵀ e, keeping joint limits; with it the mean error "
                       f"is {vals[2]:.4f} mm and {F['l1_lm_within_1mm_1deg']:.0f} % of 2,613 test poses are within 1 mm & 1°, "
                       "from warm and cold (zero-pose) starts, with 0 % self-collision and 0 % joint-limit violations."))
    cd = CategoryChartData()
    cd.categories = ["Original plan · network alone", "Our network alone", "Network + LM solver"]
    cd.add_series("Mean position error (mm)", vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(M), Inches(2.2), Inches(7.6), Inches(4.5), cd)
    ch = gf.chart
    style_chart(ch, legend=False, size=13)
    log_axis(ch.value_axis, 0.0001, 1000)
    ch.value_axis.crosses = XL_AXIS_CROSSES.MINIMUM
    ch.has_title = False
    ch.value_axis.tick_labels.number_format = '[>=1]0" mm";0.####" mm"'
    ch.value_axis.tick_labels.number_format_is_linked = False
    plot = ch.plots[0]
    plot.gap_width = 70
    ser = plot.series[0]
    for i, (v, col) in enumerate(zip(vals, ["5b6573", "3a6fbf", A.C_L1])):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(col)
        dl = pt.data_label
        dl.has_text_frame = True
        dl.text_frame.text = f"{A.fmt_mm(v)} mm"
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
        r_ = dl.text_frame.paragraphs[0].runs[0]
        r_.font.size, r_.font.name, r_.font.color.rgb = Pt(15), FONT_SB, rgb(A.TXT)
    x = M + 8.2
    stat(s, x, 2.25, 3.9, f"{math.floor(vals[0] / vals[1] * 10) / 10:g}×", "better network than the original plan")
    stat(s, x, 3.75, 3.9, f"{F['l1_lm_within_1mm_1deg']:.0f} %", "within 1 mm & 1° after the LM solver")
    stat(s, x, 5.25, 3.9, "0 %", "self-collisions · joint-limit violations")

    # 7 ---------------------------------------------------------------------------------------- self-collision
    lm_meta = A._json(A.WS / "library" / "runs" / A.L1_LM_RUN / "meta.json")
    cm = lm_meta["collision_metrics"]["test"]
    s = d.slide("Never colliding with itself", "LEVEL 1 · SAFETY", A.C_L1,
                "A learned collision classifier guards both the IK network and the Level 2 agent",
                notes=(f"Collision classifier C(q): {cm['accuracy']:.1f} % accuracy, precision {cm['precision']:.1f} %, "
                       f"recall {cm['recall']:.1f} % on {cm['n']:,} unseen test poses. It is used twice: in the Level 1 "
                       "training loss (loss = MSE + 0.05·σ(C(q̂)) pushes predictions away from self-colliding "
                       "configurations) and in the Level 2 reward (− 0.3·σ(C(θ))). The final IK answer is additionally "
                       "checked exactly in PyBullet – 0 % self-collisions on the test set."))
    for i, (v, lab) in enumerate(((cm["accuracy"], "accuracy"), (cm["precision"], "precision"), (cm["recall"], "recall"))):
        stat(s, M + i * 2.75, 2.3, 2.5, f"{v:.1f} %", lab, value_size=34)
    text(s, M, 3.85, 8, 0.3, f"on {cm['n']:,} unseen test poses", 12, A.MUTED)
    box(s, M, 4.4, 8.0, 1.9, fill=A.SURF)
    text(s, M + 0.35, 4.6, 7.4, 0.35, "IN THE LEVEL 1 TRAINING LOSS", 11, A.MUTED, bold=True)
    text(s, M + 0.35, 5.05, 7.4, 0.7, "loss = MSE + 0.05 · σ(C(q̂))", 28, A.TXT, font="Cambria Math")
    text(s, M + 0.35, 5.8, 7.4, 0.4, "and in the Level 2 reward: − 0.3 · σ(C(θ))", 14, A.TXT2)
    flow = [("Predicted joints", A.MUTED), ("Classifier C", A.C_L1), ("Penalty", A.C_COLL)]
    for i, (lab, col) in enumerate(flow):
        y = 2.3 + i * 1.4
        box(s, M + 8.6, y, 3.5, 0.95, fill=A.SURF2)
        dot(s, M + 8.95, y + 0.475, 0.11, col)
        text(s, M + 9.25, y + 0.3, 2.8, 0.4, lab, 16, A.TXT, bold=True)
        if i < 2:
            a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(M + 10.2), Inches(y + 1.0), Inches(0.3), Inches(0.38))
            a.fill.solid()
            a.fill.fore_color.rgb = rgb(A.BORDER)
            a.line.fill.background()
            _no_shadow(a)

    # 8 ---------------------------------------------------------------------------------------- level 2 idea
    o, h1, _, act = F["l2_layers"]
    s = d.slide("Level 2: a tiny agent that steers in real time", "LEVEL 2 · STEER", A.C_L2,
                notes=("Residual RL: the command is the potential-field controller plus 2 × the policy's correction, "
                       "clipped. With a zero correction the agent is the potential field, so it only learns when and how "
                       "to deviate – mainly how to escape the potential field's local minima. Its observation includes, "
                       "per obstacle, the kinematics-aware escape direction w·Jᵀn (Jacobian at the closest robot point × "
                       f"obstacle normal); {o} inputs in total, {act} joint corrections out. Trained with PPO, 20 M steps. "
                       "Picture: same mission – the solid arm is the agent, the translucent green arm the straight path, "
                       "which collides."))
    if "steer" in img:
        picture(s, img["steer"], M, 1.75, w=7.6, h=7.6 * 1012 / 1920)
    x = M + 8.1
    box(s, x, 1.75, 4.0, 1.55, fill=A.SURF)
    text(s, x + 0.3, 1.92, 3.5, 0.3, "RESIDUAL RL", 11, A.MUTED, bold=True)
    text(s, x + 0.3, 2.3, 3.6, 0.6, "u = u_pf + 2 · a", 28, A.TXT, font="Cambria Math")
    text(s, x + 0.3, 2.85, 3.6, 0.35, "potential field + learned correction", 12, A.TXT2)
    for i, (big, lab) in enumerate(((f"{o}", "inputs: joints, goal, obstacles, escape directions"),
                                    (f"2 × {h1}", "hidden units (MLP)"), (f"{act}", "joint corrections out"))):
        y = 3.55 + i * 0.95
        box(s, x, y, 4.0, 0.8, fill=A.SURF)
        text(s, x + 0.3, y + 0.14, 1.3, 0.5, big, 22, A.TXT, bold=True)
        text(s, x + 1.55, y + 0.2, 2.35, 0.5, lab, 12, A.TXT2)

    # 9 ---------------------------------------------------------------------------------------- disturbances
    s = d.slide("What can go wrong – and does", "LEVEL 2 · SCENARIOS", A.C_L2,
                "Randomised in every training and test episode",
                notes=("Sudden obstacle: a sphere appears on the path mid-move. Drifting obstacles: spheres move while the "
                       "arm moves. Joint slip: one joint jumps by a random amount (hardware position error). Gain error: "
                       "every actuator executes its command with a random gain / bias error. 'Everything at once' combines "
                       "them. Obstacles are spheres; the arm is moved kinematically."))
    items = [("Sudden obstacle", "appears on the path mid-move", A.C_STUCK),
             ("Drifting obstacles", "move while the arm moves", A.C_COLL),
             ("Joint slip", "a joint jumps – hardware position error", A.C_L2),
             ("Gain error", "actuators do not do exactly what they are told", A.C_L1)]
    for i, (head, sub, col) in enumerate(items):
        x = M + i * 3.07
        box(s, x, 2.4, 2.85, 3.6, fill=A.SURF)
        cx, cy = x + 1.42, 3.45
        if i == 0:                                                     # sphere + burst ring
            box(s, cx - 0.62, cy - 0.62, 1.24, 1.24, fill=None, line=col, shape=MSO_SHAPE.OVAL)
            dot(s, cx, cy, 0.38, col)
        elif i == 1:                                                   # sphere + motion arrow
            dot(s, cx - 0.35, cy, 0.38, col)
            a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(cx + 0.1), Inches(cy - 0.15), Inches(0.75), Inches(0.3))
            a.fill.solid()
            a.fill.fore_color.rgb = rgb(A.TXT2)
            a.line.fill.background()
            _no_shadow(a)
        elif i == 2:                                                   # two links, one rotated
            box(s, cx - 0.75, cy - 0.12, 0.8, 0.24, fill=A.TXT2, line=None, radius=0.5)
            l2 = box(s, cx - 0.05, cy - 0.12, 0.8, 0.24, fill=col, line=None, radius=0.5)
            l2.rotation = -28
            dot(s, cx - 0.02, cy, 0.14, A.SURF2)
        else:                                                          # dial with an off-centre needle
            box(s, cx - 0.55, cy - 0.55, 1.1, 1.1, fill=None, line=col, shape=MSO_SHAPE.OVAL)
            n_ = box(s, cx - 0.04, cy - 0.45, 0.08, 0.47, fill=A.TXT2, line=None, radius=0.5)
            n_.rotation = 35
            dot(s, cx, cy, 0.08, col)
        text(s, x + 0.3, 4.55, 2.4, 0.4, head, 18, A.TXT, bold=True)
        text(s, x + 0.3, 5.0, 2.4, 0.8, sub, 13, A.TXT2)

    # 10 --------------------------------------------------------------------------------------- results chart
    scen = [("combined", "Everything at once"), ("static", "Obstacles from the start"), ("mixed", "Random mix")]
    s = d.slide(f"Collisions {comb['baseline']['collision_pct']:.0f} % → {comb['policy']['collision_pct']:.0f} %",
                "LEVEL 2 · RESULTS", A.C_L2, f"Collision rate on {F['l2_n']} unseen episodes per scenario, identical seeds",
                notes=("Same 200 unseen episodes and seeds for every controller. Success / collisions, everything at once: "
                       f"agent {comb['policy']['success_pct']:.0f} % / {comb['policy']['collision_pct']:.1f} %, potential field "
                       f"{comb['apf']['success_pct']:.0f} % / {comb['apf']['collision_pct']:.1f} %, straight path "
                       f"{comb['baseline']['success_pct']:.0f} % / {comb['baseline']['collision_pct']:.0f} %. The agent's "
                       "remaining failures are mostly 'stuck' (timeouts), not collisions. Honest note: on drifting "
                       "obstacles alone the potential field is as good (95.5 % vs 94.5 %)."))
    cd = CategoryChartData()
    cd.categories = [lab for _, lab in scen]
    for key, lab in (("policy", "DRL agent"), ("apf", "Potential field"), ("baseline", "Straight path")):
        cd.add_series(lab, [ev[k][key]["collision_pct"] for k, _ in scen])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(M), Inches(2.2), Inches(8.0), Inches(4.6), cd)
    ch = gf.chart
    style_chart(ch, legend=True, size=13)
    ch.value_axis.maximum_scale, ch.value_axis.minimum_scale = 60, 0
    ch.value_axis.tick_labels.number_format = '0" %"'
    ch.value_axis.tick_labels.number_format_is_linked = False
    plot = ch.plots[0]
    plot.gap_width, plot.overlap = 60, -5
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '0.0" %"', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.color.rgb, dl.font.name = Pt(11), rgb(A.TXT), FONT
    for ser, col in zip(plot.series, (A.C_AGENT, A.C_APF, A.C_STRAIGHT)):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(col)
    x = M + 8.4
    text(s, x, 2.3, 3.7, 0.3, "SUCCESS · EVERYTHING AT ONCE", 11, A.MUTED, bold=True)
    for i, (key, lab, col) in enumerate((("policy", "DRL agent", A.C_AGENT), ("apf", "Potential field", A.C_APF),
                                         ("baseline", "Straight path", A.C_STRAIGHT))):
        y = 2.75 + i * 1.05
        box(s, x, y, 3.7, 0.85, fill=A.SURF)
        box(s, x + 0.25, y + 0.33, 0.2, 0.2, fill=col, line=None, radius=0.2)
        text(s, x + 0.6, y + 0.25, 1.9, 0.4, lab, 14, A.TXT2)
        text(s, x + 2.3, y + 0.12, 1.2, 0.6, f"{comb[key]['success_pct']:.0f} %", 26, A.TXT, bold=True, align=PP_ALIGN.RIGHT)
    text(s, x, 5.95, 3.7, 0.7, f"Stuck: agent {comb['policy']['timeout_pct']:.0f} % · potential field "
         f"{comb['apf']['timeout_pct']:.1f} %", 12, A.MUTED)

    # 11 --------------------------------------------------------------------------------------- tiny
    s = d.slide("Small enough to fly", "LEVEL 2 · FOOTPRINT", A.C_L2,
                "The trained agent needs one CPU core – the GPU is for training only",
                notes=(f"Policy: {F['l2_params']:,} parameters (2 × 64 MLP). Decision time {F['l2_us']:.1f} µs per decision "
                       f"in plain NumPy on one CPU core (≈ {round(1e6 / F['l2_us'], -3):,.0f} decisions per second); no "
                       "PyTorch and no GPU at run time. Training used the GPU for the PPO updates and 80 simulated arms "
                       f"on 10 CPU workers (20 M steps, {F['l2_minutes']:.0f} min). For scale: the Level 1 network has "
                       f"{F['l1_net_params'] / 1e6:.2f} M parameters."))
    big = [(f"{F['l2_params']:,}", "parameters", "2 × 64 MLP"),
           (f"{F['l2_us']:.1f} µs", "per decision", "NumPy · one CPU core"),
           ("No GPU", "needed on board", "GPU only for training")]
    for i, (v, lab, sub) in enumerate(big):
        x = M + i * 4.1
        box(s, x, 2.5, 3.85, 3.0, fill=A.SURF)
        text(s, x + 0.35, 2.85, 3.3, 1.1, v, 50, A.TXT, bold=True)
        text(s, x + 0.37, 4.05, 3.3, 0.4, lab, 18, A.TXT2)
        text(s, x + 0.37, 4.6, 3.3, 0.4, sub, 13, A.MUTED)
    text(s, M, 5.95, 12, 0.4, f"≈ {round(1e6 / F['l2_us'], -3):,.0f} decisions per second", 18, A.TXT2, bold=True)

    # 12 --------------------------------------------------------------------------------------- mission
    s = d.slide("Mission: solve + steer, end to end", "MISSION · L1 + L2", A.C_MISSION,
                f"{md['n']} dynamic missions (drifting + sudden obstacles + joint slip), xArm6",
                notes=(f"{md['n']} missions, dynamic scenario, identical seeds per controller. IK solved in "
                       f"{md['ik_success_pct']:.0f} % (median error {md['ik_pos_err_mm']:.3f} mm); the straight path was "
                       f"blocked in {md['predicted_conflict_pct']:.0f} %; the agent reached the target in "
                       f"{pol['success_pct']:.0f} % (collisions {pol['collision_pct']:.0f} %, stuck {pol['timeout_pct']:.0f} %); "
                       f"0 % self-collisions; median final gripper error {md['median_final_mm']:.3f} mm. Screenshot: the "
                       "Studio's Mission page with the five pipeline steps for one mission."))
    if "mission_ui" in img:
        picture(s, img["mission_ui"], M, 2.25, w=7.4, h=7.4 * 1080 / 1920)
    tiles = [(f"{md['ik_success_pct']:.0f} %", "IK solved"), (f"{pol['success_pct']:.0f} %", "reached the target"),
             (f"{pol['self_collision_pct']:.0f} %", "self-collisions"), (f"{md['median_final_mm']:.3f} mm", "median final error")]
    for i, (v, lab) in enumerate(tiles):
        x = M + 7.85 + (i % 2) * 2.12
        y = 2.25 + (i // 2) * 2.12
        box(s, x, y, 1.97, 1.92, fill=A.SURF)
        text(s, x + 0.22, y + 0.35, 1.7, 0.7, v, 26 if "mm" in v else 32, A.TXT, bold=True)
        text(s, x + 0.24, y + 1.15, 1.6, 0.6, lab, 13, A.TXT2)

    # 13 --------------------------------------------------------------------------------------- iterations
    s = d.slide("How we got here", "LEVEL 2 · ITERATIONS", A.C_L2,
                "Success and collisions on “everything at once” across five design iterations",
                notes=("Five design iterations of the Level 2 agent, evaluated on 'everything at once' (success / "
                       "collisions): 30/70 → 37/55 → 67/3.5 → 78/2.5 → 84/2.0. Iteration 4 is the 5 M-step agent "
                       "(observation v1, 87 inputs); iteration 5 the 20 M-step agent with observation v2 (101 inputs: "
                       "obstacle velocities, closing speed, a 'stuck' counter and elapsed time).\n\n[TEAM: add the design "
                       "change behind iterations 1–3 – they are not recorded in the run files.]"))
    cd = CategoryChartData()
    cd.categories = [f"Iteration {i + 1}" for i in range(len(DESIGN_ITER))]
    cd.add_series("Success", [a for a, _ in DESIGN_ITER])
    cd.add_series("Collisions", [b for _, b in DESIGN_ITER])
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(M), Inches(2.2), Inches(12.1), Inches(4.5), cd)
    ch = gf.chart
    style_chart(ch, legend=True, size=13)
    ch.value_axis.maximum_scale, ch.value_axis.minimum_scale, ch.value_axis.major_unit = 100, 0, 25
    ch.value_axis.tick_labels.number_format = '0" %"'
    ch.value_axis.tick_labels.number_format_is_linked = False
    plot = ch.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = 'General" %"', False
    dl.position = XL_LABEL_POSITION.ABOVE
    dl.font.size, dl.font.color.rgb, dl.font.name = Pt(13), rgb(A.TXT), FONT_SB
    for ser, col in zip(plot.series, (A.C_REACHED, A.C_COLL)):
        ser.smooth = False
        ser.format.line.color.rgb = rgb(col)
        ser.format.line.width = Pt(3)
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = 9
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = rgb(col)
        ser.marker.format.line.color.rgb = rgb(col)

    # 14 --------------------------------------------------------------------------------------- limits
    s = d.slide("Honest limits & next steps", "OUTLOOK", A.MUTED,
                notes=(f"Limits: in 'everything at once' the agent gets stuck in {comb['policy']['timeout_pct']:.0f} % of "
                       "episodes (it does not collide, it times out); all results are simulation (PyBullet), the arm is "
                       "moved kinematically and obstacles are spheres; the Level 2 agent is trained for the xArm6 only.\n\n"
                       "Next: when stuck, fall back to another IK branch (a different joint solution for the same pose); "
                       "train a KUKA agent; load a real space-manipulator URDF; add dynamics (torques, a free-floating "
                       "satellite base).\n\nQ&A."))
    cols = [("Limits", A.C_STUCK, [f"Stuck in {comb['policy']['timeout_pct']:.0f} % of “everything at once”",
                                   "Simulation only – kinematic arm",
                                   "Obstacles are spheres",
                                   "Level 2 agent trained on the xArm6"]),
            ("Next steps", A.C_MISSION, ["Fall back to another IK branch when stuck",
                                         "Train a KUKA agent",
                                         "Real space-arm URDF",
                                         "Dynamics: torques, floating base"])]
    for i, (head, col, items_) in enumerate(cols):
        x = M + i * 6.17
        box(s, x, 1.85, 5.9, 4.75, fill=A.SURF)
        dot(s, x + 0.45, 2.35, 0.11, col)
        text(s, x + 0.7, 2.17, 4.8, 0.45, head, 22, A.TXT, bold=True)
        for j, it in enumerate(items_):
            y = 3.0 + j * 0.85
            box(s, x + 0.35, y, 5.2, 0.68, fill=A.SURF2, radius=0.15)
            text(s, x + 0.6, y + 0.18, 4.8, 0.4, it, 15, A.TXT2)
    return d


def main():
    F = A.load_facts()
    img = prep_images(F)
    missing = [k for k in ("robot_xarm6", "robot_kuka_iiwa", "robot_panda", "browse", "steer", "mission_ui") if k not in img]
    if missing:
        print("missing screenshots (run capture_studio.py):", ", ".join(missing))
    d = build(F, img)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    d.prs.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({d.n} slides)")


if __name__ == "__main__":
    main()
